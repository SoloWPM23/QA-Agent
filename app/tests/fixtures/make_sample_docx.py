"""Generate sample_test_suite.docx reproducing the reference template shape.

Table 0: reference tipe (rows: Tipe, Status Code, Schema, JSONPath, Regex,
Contains) -> non test case, must be skipped.
Table 1: populated TC-001 (14 rows).
Table 2: empty TC-00X template (14 rows) -> case_id fails pattern, needs_review.

Run from project root: venv/Scripts/python.exe -m app.tests.fixtures.make_sample_docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document


def _make_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    for i, (label, value) in enumerate(rows):
        table.cell(i, 0).text = label
        table.cell(i, 1).text = value


def main(out_path: Path) -> None:
    doc = Document()

    # Table 0: reference (non test case).
    _make_table(
        doc,
        [
            ["Tipe", "Contoh"],
            ["Status Code", "200, 201, 404"],
            ["Schema", "JSON Schema"],
            ["JSONPath", "$.data.id"],
            ["Regex", "^TC"],
            ["Contains", "abc"],
        ],
    )

    # Table 1: populated TC-001.
    _make_table(
        doc,
        [
            ["Test Case TC-001", "Test Case TC-001"],
            ["ID", "TC-001"],
            ["Judul", "Ambil daftar pengguna"],
            ["Deskripsi", "Memastikan daftar pengguna dapat diambil"],
            ["Method", "GET"],
            ["Path", "/api/v1/users"],
            ["Headers", "Authorization: Bearer xxx"],
            ["Query Params", "page: 1"],
            ["Body (JSON)", ""],
            ["Expected Status Code", "200"],
            ["Expected Schema", ""],
            ["JSONPath Checks", "$.page"],  # single -> bucket list
            ["Regex", ""],
            ["Contains", "users"],
        ],
    )

    # Table 2: empty template TC-00X -> needs_review.
    _make_table(
        doc,
        [
            ["Test Case TC-00X", "Test Case TC-00X"],
            ["ID", "TC-00X"],
            ["Judul", ""],
            ["Deskripsi", ""],
            ["Method", ""],
            ["Path", ""],
            ["Headers", ""],
            ["Query Params", ""],
            ["Body (JSON)", ""],
            ["Expected Status Code", ""],
            ["Expected Schema", ""],
            ["JSONPath Checks", ""],
            ["Regex", ""],
            ["Contains", ""],
        ],
    )

    doc.save(str(out_path))
    print(f"fixture dibuat: {out_path}")


if __name__ == "__main__":
    here = Path(__file__).resolve().parent
    main(here / "sample_test_suite.docx")
